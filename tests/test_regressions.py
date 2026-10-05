import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from enfermeria_app.database import Database, normalize_blood_type, normalize_formation_stage
from enfermeria_app.excel_io import as_date, export_inventory, export_seminarians, import_inventory, import_seminarians, infer_location, normalize_product_name
from enfermeria_app.pdf_export import export_all_dossiers_pdf, export_control_pdf, export_emergency_summary_pdf, export_seminarian_dossier_pdf
from enfermeria_app.panel_data import read_panel_snapshot
from enfermeria_app.notifier import notification_signature, notification_text


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.db = Database(self.root / "enfermeria.db")

    def item(self, name="Paracetamol", kind="medicine", minimum=5):
        return self.db.save_item({"name":name,"type":kind,"minimum_stock":minimum,"location":"DA","unit":"unidad"})

    def seminarian(self, external_id="0100000001", first="Juan", last="Pérez", allergies=""):
        return self.db.insert_seminarian({"external_id":external_id,"first_name":first,"last_name":last,"allergies":allergies,"consent":1})

    def test_expired_lot_is_never_delivered_but_can_be_discarded(self):
        item=self.item(); sem=self.seminarian()
        expired=(date.today()-timedelta(days=1)).isoformat(); valid=(date.today()+timedelta(days=30)).isoformat()
        self.db.add_stock(item,5,"VENCIDO",expired); self.db.add_stock(item,5,"VIGENTE",valid)
        movement=self.db.record_output(item,3,"administered",sem,"1 tableta","")
        lots={r["lot_number"]:r["quantity"] for r in self.db.query("SELECT lot_number,quantity FROM lots WHERE item_id=?",(item,))}
        self.assertEqual(lots,{"VENCIDO":5,"VIGENTE":2})
        with self.assertRaisesRegex(ValueError,"sin contar lotes vencidos"):
            self.db.record_output(item,3,"delivery",sem)
        self.db.record_output(item,2,"discard",None,reason="Vencido")
        lots={r["lot_number"]:r["quantity"] for r in self.db.query("SELECT lot_number,quantity FROM lots WHERE item_id=?",(item,))}
        self.assertEqual(lots["VENCIDO"],3)
        self.db.void_movement(movement,"Registro de prueba")
        lots={r["lot_number"]:r["quantity"] for r in self.db.query("SELECT lot_number,quantity FROM lots WHERE item_id=?",(item,))}
        self.assertEqual(lots["VIGENTE"],5)

    def test_lot_expiring_today_is_valid_all_day(self):
        item=self.item(); sem=self.seminarian(); self.db.add_stock(item,2,"HOY",date.today().isoformat())
        self.db.record_output(item,1,"delivery",sem)
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0]["stock"],1)
        self.assertEqual(self.db.stats()["expired"],0)

    def test_zero_minimum_is_preserved(self):
        item=self.item(minimum=0)
        self.assertEqual(self.db.query("SELECT minimum_stock FROM items WHERE id=?",(item,))[0][0],0)

    def test_partial_reimport_does_not_erase_health_data(self):
        self.db.upsert_seminarian({"external_id":"0123456789","first_name":"Ana","last_name":"López","birth_date":"2000-01-01","allergies":"Penicilina","room":"12"})
        wb=Workbook(); ws=wb.active; ws.title="Fichas"; ws.append(["Número de identificación","Nombres","Apellidos","Fecha de nacimiento","Año de formación"]); ws.append([123456789,"Ana","López","2000-01-01","Configuración 2"])
        path=self.root/"fichas.xlsx"; wb.save(path); result=import_seminarians(self.db,path)
        self.assertEqual(result["updated"],1)
        row=self.db.query("SELECT external_id,allergies,room,formation_year FROM seminarians")[0]
        self.assertEqual((row["external_id"],row["allergies"],row["room"]),("0123456789","Penicilina","12"))
        self.assertEqual(row["formation_year"],"Configuración 2")

    def test_same_name_without_unique_data_requires_review(self):
        self.db.insert_seminarian({"first_name":"Juan","last_name":"Pérez","allergies":"Penicilina"})
        with self.assertRaisesRegex(ValueError,"REVISAR_DUPLICADO"):
            self.db.upsert_seminarian({"first_name":"Juan","last_name":"Pérez","allergies":"Ninguna"})
        self.assertEqual(self.db.query("SELECT allergies FROM seminarians")[0][0],"Penicilina")

    def test_equipment_loan_is_atomic(self):
        item=self.item("Tensiómetro","equipment"); sem=self.seminarian(); self.db.add_stock(item,2)
        with self.db.connect() as conn:
            conn.execute("CREATE TRIGGER fail_loan BEFORE INSERT ON equipment_loans BEGIN SELECT RAISE(ABORT,'fallo simulado'); END")
        with self.assertRaises(Exception):
            self.db.create_equipment_loan(item,sem,1,date.today().isoformat())
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],2)
        self.assertEqual(self.db.query("SELECT COUNT(*) FROM movements WHERE type='loan'")[0][0],0)

    def test_equipment_loan_and_return_share_consistent_stock(self):
        item=self.item("Nebulizador","equipment"); sem=self.seminarian(); self.db.add_stock(item,1)
        loan=self.db.create_equipment_loan(item,sem,1,(date.today()+timedelta(days=3)).isoformat())
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],0)
        self.db.return_equipment_loan(loan)
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],1)
        row=self.db.query("SELECT movement_id,return_movement_id,returned_at FROM equipment_loans WHERE id=?",(loan,))[0]
        self.assertTrue(row["movement_id"] and row["return_movement_id"] and row["returned_at"])
        with self.assertRaisesRegex(ValueError,"préstamo"):
            self.db.void_movement(row["return_movement_id"],"Prueba")

    def test_quick_panel_is_read_only_and_hides_seminarian_identity(self):
        medicine=self.item("Jarabe de prueba",minimum=5);self.db.add_stock(medicine,2,"PROX",(date.today()+timedelta(days=10)).isoformat())
        equipment=self.item("Tensiómetro","equipment",minimum=0);seminarian=self.seminarian(first="Nombre Privado",last="No Mostrar")
        self.db.add_stock(equipment,1);self.db.create_equipment_loan(equipment,seminarian,1,(date.today()-timedelta(days=1)).isoformat())
        before=self.db.query("SELECT COUNT(*) FROM audit_log")[0][0]
        snapshot=read_panel_snapshot(self.db.path)
        after=self.db.query("SELECT COUNT(*) FROM audit_log")[0][0]
        self.assertEqual((snapshot.expiring,snapshot.low_stock,snapshot.active_loans,snapshot.overdue_loans),(1,2,1,1))
        rendered=" ".join(f"{a.title} {a.detail} {a.state}" for a in snapshot.alerts)
        self.assertIn("Jarabe de prueba",rendered);self.assertIn("Tensiómetro",rendered)
        self.assertNotIn("Nombre Privado",rendered);self.assertNotIn("No Mostrar",rendered)
        self.assertEqual(before,after)

    def test_linux_notification_summary_is_private_and_stable(self):
        medicine=self.item("Medicamento reservado",minimum=5);self.db.add_stock(medicine,1,"PROX",(date.today()+timedelta(days=5)).isoformat())
        equipment=self.item("Nebulizador","equipment",minimum=0);seminarian=self.seminarian(first="Identidad",last="Confidencial")
        self.db.add_stock(equipment,1);self.db.create_equipment_loan(equipment,seminarian,1,(date.today()-timedelta(days=2)).isoformat())
        snapshot=read_panel_snapshot(self.db.path);title,message=notification_text(snapshot)
        self.assertEqual(title,"Enfermería requiere atención")
        self.assertIn("stock bajo",message);self.assertIn("préstamo(s) atrasado(s)",message)
        self.assertNotIn("Identidad",message);self.assertNotIn("Confidencial",message)
        self.assertEqual(notification_signature(snapshot),notification_signature(read_panel_snapshot(self.db.path)))

    def test_permissions_backup_and_month_expiry(self):
        if os.name != "nt":
            self.assertEqual(os.stat(self.db.path).st_mode & 0o777,0o600)
            self.assertEqual(os.stat(Path(self.db.path).parent).st_mode & 0o777,0o700)
        reopened=Database(self.db.path)
        self.assertTrue(list((Path(reopened.path).parent/"backups").glob("*.db")))
        self.assertEqual(as_date("feb 2028"),"2028-02-29")

    def test_allergy_notice_is_advisory(self):
        item=self.item("Amoxicilina 500 mg"); sem=self.seminarian(allergies="Alergia a amoxicilina y penicilina")
        message,match=self.db.medication_safety_notice(item,sem)
        self.assertTrue(match); self.assertIn("amoxicilina",message.lower())

    def test_import_name_normalization(self):
        self.assertEqual(normalize_product_name("IBUPROFENO"),"Ibuprofeno")
        self.assertEqual(normalize_product_name("  IbUproFeno   400MG  "),"Ibuprofeno 400mg")
        self.assertEqual(normalize_product_name("insulina nph 100 ui"),"Insulina NPH 100 UI")

    def test_inventory_import_groups_duplicate_product_rows_as_lots(self):
        wb=Workbook();ws=wb.active;ws.title="Inventario"
        ws.append(["Producto","Presentación","Lote","Ubicación","Stock_Actual","Stock_Mínimo","Fecha_Vencimiento"])
        ws.append(["XAROBAN — Rivaroxabán 15 mg","Caja","74675","AC",20,5,"2028-01-31"])
        ws.append(["XAROBAN — Rivaroxabán 15 mg","Caja","76300","AC",10,5,"2028-06-30"])
        path=self.root/"inventario_lotes.xlsx";wb.save(path)
        result=import_inventory(self.db,path)
        self.assertEqual((result["created"],result["updated"],result["recognized"],result["stock_entries"]),(1,0,2,2))
        item=self.db.query("SELECT id,stock FROM items")[0]
        self.assertEqual(item["stock"],30)
        self.assertEqual([(r["lot_number"],r["quantity"],r["expires_on"]) for r in self.db.query("SELECT lot_number,quantity,expires_on FROM lots ORDER BY id")],[("74675",20,"2028-01-31"),("76300",10,"2028-06-30")])
        second=import_inventory(self.db,path)
        self.assertEqual((second["created"],second["updated"],second["stock_entries"]),(0,1,0))
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item["id"],))[0][0],30)

    def test_inventory_location_is_inferred_only_when_unassigned(self):
        self.assertEqual(infer_location("AINEs / Analgésicos"),"DA")
        self.assertEqual(infer_location("Anticoagulantes / Inhibidores del factor Xa"),"AC")
        wb=Workbook();ws=wb.active;ws.title="Inventario";ws.append(["Producto","Categoria","Ubicacion","Stock_Actual"])
        ws.append(["Ibuprofeno","AINEs / Analgésicos","Por asignar",1]);ws.append(["Amoxicilina","Antibióticos","QR",1])
        path=self.root/"ubicaciones.xlsx";wb.save(path);import_inventory(self.db,path)
        locations={r["name"]:r["location"] for r in self.db.query("SELECT name,location FROM items")}
        self.assertEqual(locations,{"Ibuprofeno":"DA","Amoxicilina":"CU"})

    def test_inventory_import_rejects_expired_lots_but_keeps_valid_ones(self):
        yesterday=(date.today()-timedelta(days=1)).isoformat();today=date.today().isoformat();future=(date.today()+timedelta(days=40)).isoformat()
        wb=Workbook();ws=wb.active;ws.title="Inventario";ws.append(["Producto","Lote","Stock_Actual","Fecha_Vencimiento"])
        ws.append(["Producto mixto","VENCIDO",4,yesterday]);ws.append(["Producto mixto","HOY",3,today]);ws.append(["Producto mixto","VIGENTE",5,future])
        ws.append(["Solo vencido","OLD",2,yesterday]);ws.append(["Fecha dudosa","X",1,"fecha incorrecta"])
        path=self.root/"vencidos.xlsx";wb.save(path);result=import_inventory(self.db,path)
        self.assertEqual((result["recognized"],result["expired_rejected"],result["expired_units"],result["errors"]),(5,2,6,1))
        self.assertEqual((result["created"],result["stock_entries"]),(1,2))
        item=self.db.query("SELECT id,stock FROM items WHERE name='Producto mixto'")[0];self.assertEqual(item["stock"],8)
        self.assertFalse(self.db.query("SELECT 1 FROM items WHERE name IN ('Solo vencido','Fecha dudosa')"))
        self.assertEqual({r["lot_number"] for r in self.db.query("SELECT lot_number FROM lots WHERE item_id=?",(item["id"],))},{"HOY","VIGENTE"})

    def test_retire_all_expired_lots_preserves_valid_stock_and_history(self):
        item=self.item();expired=(date.today()-timedelta(days=2)).isoformat();valid=(date.today()+timedelta(days=20)).isoformat()
        self.db.add_stock(item,2,"OLD-1",expired);self.db.add_stock(item,3,"OLD-2",expired);self.db.add_stock(item,4,"OK",valid)
        result=self.db.discard_expired_lots(item)
        self.assertEqual(result,{"lots":2,"quantity":5});self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],4)
        lots={r["lot_number"]:r["quantity"] for r in self.db.query("SELECT lot_number,quantity FROM lots WHERE item_id=?",(item,))}
        self.assertEqual(lots,{"OLD-1":0,"OLD-2":0,"OK":4})
        self.assertEqual(self.db.query("SELECT COUNT(*) FROM movements WHERE item_id=? AND type='discard'",(item,))[0][0],2)

    def test_manual_text_is_formatted_but_email_is_untouched(self):
        item=self.db.save_item({"name":"IBuPROFENO","type":"medicine","minimum_stock":5,"location":"DA","unit":"tableta","usage":"dolor y fiebre"})
        product=self.db.query("SELECT name,usage FROM items WHERE id=?",(item,))[0]
        self.assertEqual((product["name"],product["usage"]),("Ibuprofeno","Dolor y fiebre"))
        sem=self.db.insert_seminarian({"external_id":"0100000077","first_name":"juan carlos","last_name":"PÉREZ DE LA cruz","email":"MiCorreo+Salud@Ejemplo.COM","city":"guayaquil","consent":1})
        row=self.db.query("SELECT first_name,last_name,email,city FROM seminarians WHERE id=?",(sem,))[0]
        self.assertEqual((row["first_name"],row["last_name"],row["city"]),("Juan Carlos","Pérez de la Cruz","Guayaquil"))
        self.assertEqual(row["email"],"MiCorreo+Salud@Ejemplo.COM")

    def test_seminarian_selectors_normalize_imported_values(self):
        self.assertEqual(normalize_blood_type("o positivo"),"O+")
        self.assertEqual(normalize_blood_type("AB -"),"AB-")
        self.assertEqual(normalize_blood_type("A−"),"A-")
        self.assertEqual(normalize_blood_type("No consta"),"")
        self.assertEqual(normalize_formation_stage("propedeutico"),"Propedéutico")
        self.assertEqual(normalize_formation_stage("1 filosofia"),"1° Filosofía")
        self.assertEqual(normalize_formation_stage("4° teología"),"4° Teología")
        self.assertEqual(normalize_formation_stage("3°"),"3°")
        sem=self.db.insert_seminarian({"first_name":"Luis","last_name":"Prueba","formation_year":"2 teologia","blood_type":"a positivo"})
        row=self.db.query("SELECT formation_year,blood_type FROM seminarians WHERE id=?",(sem,))[0]
        self.assertEqual((row["formation_year"],row["blood_type"]),("2° Teología","A+"))

    def test_inventory_export_contains_no_health_sheets(self):
        from openpyxl import load_workbook
        self.item(); path=self.root/"inventario.xlsx"; export_inventory(self.db,path)
        self.assertEqual(load_workbook(path,read_only=True).sheetnames,["Inventario"])

    def test_settings_pin_audit_and_restore(self):
        self.db.set_setting("manager_name","Encargado de prueba")
        self.assertEqual(self.db.get_setting("manager_name"),"Encargado de prueba")
        self.db.set_pin("2468");self.assertTrue(self.db.verify_pin("2468"));self.assertFalse(self.db.verify_pin("1234"))
        item=self.item("Producto antes de copia");backup=self.db.create_backup(self.root/"respaldo.db")
        self.db.safe_delete_item(item)
        self.assertFalse(self.db.query("SELECT id FROM items WHERE id=?",(item,)))
        self.db.restore_backup(backup)
        self.assertTrue(self.db.query("SELECT id FROM items WHERE id=?",(item,)))
        self.assertTrue(self.db.query("SELECT COUNT(*) FROM audit_log")[0][0]>0)

    def test_purge_removes_test_data_and_backups_but_preserves_configuration(self):
        self.db.set_setting("clinic_name","Enfermería de prueba");self.db.set_pin("2468")
        item=self.item();sem=self.seminarian();self.db.add_stock(item,2,"TEST",(date.today()+timedelta(days=30)).isoformat());self.db.record_output(item,1,"delivery",sem)
        backup=self.db.create_backup();self.assertTrue(backup.exists())
        self.db.purge_all_data(delete_local_backups=True)
        for table in ("items","lots","movements","seminarians","medical_controls","equipment_loans","audit_log"):
            self.assertEqual(self.db.query(f"SELECT COUNT(*) FROM {table}")[0][0],0)
        self.assertEqual(self.db.get_setting("clinic_name"),"Enfermería de prueba");self.assertTrue(self.db.verify_pin("2468"))
        self.assertEqual(self.db.list_backups(),[])

    def test_full_lot_discard_is_audited_and_reversible(self):
        item=self.item();self.db.add_stock(item,4,"L-RET",(date.today()-timedelta(days=10)).isoformat())
        lot=self.db.query("SELECT id FROM lots WHERE item_id=?",(item,))[0][0]
        self.db.discard_lot(lot,"Producto vencido")
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],0)
        movement=self.db.query("SELECT id FROM movements WHERE item_id=? AND type='discard' ORDER BY id DESC",(item,))[0][0]
        self.db.void_movement(movement,"Retiro registrado por error")
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],4)
        self.assertEqual(self.db.query("SELECT quantity FROM lots WHERE id=?",(lot,))[0][0],4)

    def test_archived_product_can_be_restored(self):
        item=self.item();self.db.archive_item(item)
        self.assertEqual(self.db.query("SELECT archived FROM items WHERE id=?",(item,))[0][0],1)
        self.db.restore_item(item)
        self.assertEqual(self.db.query("SELECT archived FROM items WHERE id=?",(item,))[0][0],0)

    def test_only_empty_unused_product_can_be_deleted_permanently(self):
        empty_id=self.item("Registro accidental");self.db.archive_item(empty_id)
        self.db.safe_delete_item(empty_id)
        self.assertFalse(self.db.query("SELECT 1 FROM items WHERE id=?",(empty_id,)))

        used_id=self.item("Producto con historia");self.db.add_stock(used_id,1,"L-1",None)
        self.db.archive_item(used_id)
        with self.assertRaises(ValueError):self.db.safe_delete_item(used_id)

    def test_archived_seminarian_can_be_restored(self):
        sem=self.seminarian();self.db.archive_seminarian(sem)
        self.assertEqual(self.db.query("SELECT archived FROM seminarians WHERE id=?",(sem,))[0][0],1)
        self.db.restore_seminarian(sem)
        self.assertEqual(self.db.query("SELECT archived FROM seminarians WHERE id=?",(sem,))[0][0],0)

    def test_archived_seminarian_can_be_permanently_deleted_and_history_anonymized(self):
        item=self.item();sem=self.seminarian();self.db.add_stock(item,2,"L-1",(date.today()+timedelta(days=90)).isoformat())
        self.db.record_output(item,1,"delivery",sem,reason="Atención")
        self.db.save_medical_control(sem,{"measured_at":"2026-10-02T08:00:00","weight_kg":70,"height_cm":175})
        self.db.archive_seminarian(sem);result=self.db.delete_archived_seminarian(sem)
        self.assertEqual((result["medical_controls"],result["movements"]),(1,1))
        self.assertFalse(self.db.query("SELECT 1 FROM seminarians WHERE id=?",(sem,)))
        self.assertFalse(self.db.query("SELECT 1 FROM medical_controls WHERE seminarian_id=?",(sem,)))
        movement=self.db.query("SELECT seminarian_id,quantity FROM movements WHERE type='delivery'")[0]
        self.assertEqual((movement["seminarian_id"],movement["quantity"]),(None,1))
        self.assertEqual(self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0][0],1)

    def test_seminarian_with_active_equipment_loan_cannot_be_deleted(self):
        equipment=self.item("Nebulizador","equipment");sem=self.seminarian();self.db.add_stock(equipment,1)
        self.db.create_equipment_loan(equipment,sem,1,(date.today()+timedelta(days=7)).isoformat())
        self.db.archive_seminarian(sem)
        with self.assertRaisesRegex(ValueError,"pendientes de devolución"):
            self.db.delete_archived_seminarian(sem)
        self.assertTrue(self.db.query("SELECT 1 FROM seminarians WHERE id=?",(sem,)))

    def test_product_and_initial_stock_are_created_together(self):
        item=self.db.create_item_with_initial_stock({"name":"Ibuprofeno 400 mg","type":"medicine","minimum_stock":5,"location":"DA","unit":"tableta"},12,"IBU-01",(date.today()+timedelta(days=300)).isoformat(),"Inventario inicial")
        row=self.db.query("SELECT stock FROM items WHERE id=?",(item,))[0]
        self.assertEqual(row["stock"],12)
        lot=self.db.query("SELECT quantity,lot_number FROM lots WHERE item_id=?",(item,))[0]
        self.assertEqual((lot["quantity"],lot["lot_number"]),(12,"IBU-01"))
        self.assertEqual(self.db.query("SELECT COUNT(*) FROM movements WHERE item_id=? AND type='entry'",(item,))[0][0],1)

    def test_medical_controls_keep_history_and_can_be_voided(self):
        sem=self.seminarian();first=self.db.save_medical_control(sem,{"measured_at":"2026-10-02T08:00:00","weight_kg":70,"height_cm":175,"systolic":120,"diastolic":80,"glucose_mg_dl":90,"glucose_context":"fasting"})
        second=self.db.save_medical_control(sem,{"measured_at":"2026-10-09T08:00:00","weight_kg":69.5,"height_cm":175,"systolic":118,"diastolic":78})
        rows=self.db.query("SELECT id,bmi,voided_at FROM medical_controls WHERE seminarian_id=? ORDER BY measured_at",(sem,))
        self.assertEqual([r["id"] for r in rows],[first,second]);self.assertEqual(rows[0]["bmi"],22.86)
        self.db.void_medical_control(first,"Medición incorrecta")
        self.assertTrue(self.db.query("SELECT voided_at FROM medical_controls WHERE id=?",(first,))[0][0])
        self.assertEqual(self.db.query("SELECT COUNT(*) FROM medical_controls WHERE seminarian_id=?",(sem,))[0][0],2)
        path=self.root/"fichas_y_controles.xlsx";export_seminarians(self.db,path)
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=True)
        self.assertEqual(workbook.sheetnames,["Expedientes","Controles de salud"])
        self.assertEqual(workbook["Controles de salud"].max_row,2)  # encabezado + único control vigente
        with self.assertRaisesRegex(ValueError,"anulado"):
            self.db.save_medical_control(sem,{"measured_at":"2026-10-10T08:00:00"},first)

    def test_quarantine_blocks_delivery_without_changing_stock_or_location(self):
        item=self.item();sem=self.seminarian();self.db.add_stock(item,4,"REVISION",(date.today()+timedelta(days=60)).isoformat())
        lot=self.db.query("SELECT id FROM lots WHERE item_id=?",(item,))[0][0]
        self.db.set_lot_quarantine(lot,True,"Sello roto")
        row=self.db.query("SELECT stock,location FROM items WHERE id=?",(item,))[0]
        self.assertEqual((row["stock"],row["location"]),(4,"DA"));self.assertEqual(self.db.stats()["quarantined"],1)
        with self.assertRaisesRegex(ValueError,"cuarentena"):
            self.db.record_output(item,1,"delivery",sem)
        self.db.set_lot_quarantine(lot,False);self.db.record_output(item,1,"delivery",sem)
        self.assertEqual(self.db.query("SELECT quantity FROM lots WHERE id=?",(lot,))[0][0],3)

    def test_encrypted_backup_transfer_and_supervised_pin_recovery(self):
        self.item();self.db.set_setting("recovery_supervisor_name","Rector Institucional")
        recovery=self.db.set_pin("1234")
        new_recovery=self.db.reset_pin_with_recovery(recovery,"rector institucional","5678")
        self.assertTrue(self.db.verify_pin("5678"));self.assertNotEqual(recovery,new_recovery)
        with self.assertRaisesRegex(ValueError,"no es válido"):
            self.db.reset_pin_with_recovery(recovery,"Rector Institucional","9999")
        backup=self.root/"respaldo.enfbackup";self.db.create_encrypted_backup(backup,"Contraseña fuerte 2026")
        self.db.save_item({"name":"Producto posterior","type":"medicine"})
        self.db.restore_encrypted_backup(backup,"Contraseña fuerte 2026")
        self.assertFalse(self.db.query("SELECT 1 FROM items WHERE name='Producto posterior'"))
        with self.assertRaises(ValueError):self.db.restore_encrypted_backup(backup,"Contraseña equivocada")
        transfer=self.root/"entrega.enftransfer";self.db.export_transfer_package(transfer,"Otra contraseña fuerte","3.4.0")
        self.db.set_setting("institution_name","Institución temporal")
        manifest=self.db.import_transfer_package(transfer,"Otra contraseña fuerte")
        self.assertEqual(manifest["app_version"],"3.4.0");self.assertNotEqual(self.db.get_setting("institution_name"),"Institución temporal")

    def test_emergency_pdf_uses_only_latest_valid_control(self):
        sem=self.seminarian();old=self.db.save_medical_control(sem,{"measured_at":"2026-10-01T08:00:00","weight_kg":70});self.db.void_medical_control(old,"Dato errado")
        self.db.save_medical_control(sem,{"measured_at":"2026-10-02T08:00:00","weight_kg":71})
        target=self.root/"emergencia.pdf";export_emergency_summary_pdf(self.db,sem,target,"Encargado")
        self.assertTrue(target.exists());self.assertEqual(target.read_bytes()[:4],b"%PDF");self.assertGreater(target.stat().st_size,2000)

    def test_extended_seminarian_profile_and_pdf_exports(self):
        sem=self.db.insert_seminarian({"external_id":"0100000099","first_name":"José","last_name":"Moscati","birth_date":"1880-07-25",
            "birth_city":"Benevento","birth_country":"Italia","nationality":"Italiana","diocese":"Arquidiócesis de Nápoles",
            "parish":"Santa Lucía","address":"Via de prueba","neighborhood":"Centro","city":"Nápoles","province":"Campania",
            "country":"Italia","personal_phone":"0990000000","email":"jose@example.org","blood_type":"O+","allergies":"Ninguna conocida","consent":1})
        control=self.db.save_medical_control(sem,{"measured_at":"2026-10-02T09:00:00","weight_kg":72,"height_cm":176,"systolic":120,"diastolic":80,"pulse":70,"glucose_mg_dl":92,"glucose_context":"fasting","recorded_by":"Encargado"})
        row=self.db.query("SELECT birth_city,diocese,address,personal_phone FROM seminarians WHERE id=?",(sem,))[0]
        self.assertEqual((row["birth_city"],row["diocese"],row["address"],row["personal_phone"]),("Benevento","Arquidiócesis de Nápoles","Via de prueba","0990000000"))
        targets=[self.root/"control.pdf",self.root/"ficha.pdf",self.root/"todos.pdf"]
        export_control_pdf(self.db,control,targets[0]);export_seminarian_dossier_pdf(self.db,sem,targets[1]);export_all_dossiers_pdf(self.db,targets[2])
        for target in targets:
            self.assertTrue(target.exists());self.assertGreater(target.stat().st_size,2500);self.assertEqual(target.read_bytes()[:4],b"%PDF")


if __name__ == "__main__":
    unittest.main()
