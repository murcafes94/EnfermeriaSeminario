"""Regresión del fallo real de arranque causado por un NumPy incompleto."""
import subprocess
import sys
import unittest
from pathlib import Path


class PackagedStartupTests(unittest.TestCase):
    def test_excel_roundtrip_with_incomplete_numpy(self):
        project = Path(__file__).resolve().parents[1]
        code = '''
import sys, types, runpy, tempfile
from pathlib import Path
sys.modules["numpy"] = types.ModuleType("numpy")
sys.frozen = True
runpy.run_path("run.py", run_name="packaged_entry_test")
from openpyxl import Workbook, load_workbook
assert sys.modules["numpy"] is None
with tempfile.TemporaryDirectory() as directory:
    path = Path(directory) / "numeros.xlsx"
    book = Workbook()
    book.active.append(["Stock", 15, 72.5])
    book.save(path)
    loaded = load_workbook(path)
    assert list(loaded.active.values) == [("Stock", 15, 72.5)]
    loaded.close()
'''
        result = subprocess.run([sys.executable, "-c", code], cwd=project,
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
