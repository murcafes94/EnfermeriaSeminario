from __future__ import annotations

import re
import unicodedata
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .database import Database


MEDICAL_ACRONYMS = {"IV", "IM", "SC", "SL", "NPH", "UI", "UCI", "ORS", "SR", "XR", "CR"}


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def normalize_product_name(value: Any) -> str:
    text = re.sub(r"\s+", " ", excel_text(value)).strip().lower()
    if not text:
        return ""
    text = text[0].upper() + text[1:]
    def restore_acronym(match):
        word = match.group(0)
        upper = word.upper()
        if upper in MEDICAL_ACRONYMS or (re.fullmatch(r"[a-zA-Z]{1,2}\d{1,3}", word) and len(word) <= 5):
            return upper
        return word
    return re.sub(r"\b[a-zA-Z]+\d*\b", restore_acronym, text)


def infer_location(category: Any, product_name: Any = "") -> str:
    """Sugiere la ubicación física a partir de la categoría del inventario.

    Solo se usa cuando el archivo no proporciona una ubicación concreta.
    """
    text=normalize(f"{category or ''} {product_name or ''}")
    rules=(
        ("NO",("otologic","nasal")),
        ("UR",("urinari","urolog","pde_5","acido_urico","antigotos")),
        ("DR",("dermatolog","topico","topicos","emoliente","piel","hematoma")),
        ("SH",("hipnot","sedante","opioide")),
        ("HB",("hepat","biliar")),
        ("DM",("antidiabet","diabetes","sglt2","dpp_4","biguanida")),
        ("CL",("hipolipem","estatina","fibrato","colesterol")),
        ("AC",("anticoagul","antihemorrag","hemostat","venoton","vasoprotector","factor_xa")),
        ("CV",("antihipert","antiarrit","arni","calcioantagonista")),
        ("RS",("antigripal","antitus","mucolit","expectorante","descongestionante","respiratori")),
        ("AB",("antibiot","antibacter","penicilina","cefalosporina","fluoroquinolona","sulfonamida","antiparasitari")),
        ("AL",("antihistamin","antialerg","antileucotrien")),
        ("GR",("antiulcer","antiacid","antirreflujo","reflujo","bomba_de_protones","protector_gastr")),
        ("IN",("laxante","probiot","intestinal","estrenimiento")),
        ("DG",("antiemet","procinet","antiespasmod","enzima_digest","digestivo","motilidad_gastro")),
        ("AF",("antifung",)),
        ("VH",("vitamina","hierro","antianem","mineral","suplement")),
        ("DA",("aine","analges","antipiret","antiinflam","corticoide","dolor")),
    )
    for code,keywords in rules:
        if any(keyword in text for keyword in keywords):return code
    return "OT"


def as_date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    months = {"ene":1,"feb":2,"mar":3,"abr":4,"may":5,"jun":6,"jul":7,"ago":8,"sep":9,"oct":10,"nov":11,"dic":12}
    parts = text.lower().replace("de", " ").split()
    if len(parts) >= 3 and parts[1][:3] in months:
        try:
            return f"{int(parts[2]):04d}-{months[parts[1][:3]]:02d}-{int(parts[0]):02d}"
        except ValueError:
            return None
    if len(parts) == 2 and parts[0][:3] in months:
        try:
            year = int(parts[1]); month = months[parts[0][:3]]
            return f"{year:04d}-{month:02d}-{monthrange(year, month)[1]:02d}"
        except ValueError:
            return None
    return None


def excel_text(value: Any, restore_leading_zero: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    if restore_leading_zero and text.isdigit() and len(text) == 9:
        text = "0" + text
    return text


def header_map(ws) -> dict[str, int]:
    return {normalize(cell.value): i for i, cell in enumerate(ws[1]) if cell.value not in (None, "")}


def pick(row, headers: dict[str, int], *aliases, default=""):
    for alias in aliases:
        index = headers.get(normalize(alias))
        if index is not None and index < len(row):
            value = row[index]
            if value is not None:
                return value
    return default


def split_legacy_name(full_name: Any) -> tuple[str, str]:
    parts = str(full_name or "").strip().split()
    if len(parts) <= 1:
        return (parts[0] if parts else "", "")
    # En el formulario nuevo se piden por separado. Esto solo cubre archivos antiguos.
    middle = max(1, len(parts) // 2)
    return " ".join(parts[:middle]), " ".join(parts[middle:])


def import_seminarians(db: Database, filename: str | Path) -> dict[str, int]:
    wb = load_workbook(filename, data_only=True)
    preferred = ["Respuestas de formulario 1", "Form Responses 1", "Fichas"]
    ws = next((wb[name] for name in preferred if name in wb.sheetnames), wb[wb.sheetnames[0]])
    headers = header_map(ws)
    result = {"created": 0, "updated": 0, "review": 0, "skipped": 0, "errors": 0}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        first = pick(row, headers, "Nombres", "Nombre", "first_name")
        last = pick(row, headers, "Apellidos", "last_name")
        if not first or not last:
            full = pick(row, headers, "Nombre completo", "Seminarista")
            first, last = split_legacy_name(full)
        if not first or not last:
            result["skipped"] += 1
            continue
        data = {
            "external_id": excel_text(pick(row, headers, "Número de identificación", "Cedula", "Documento", "external_id"), True),
            "first_name": excel_text(first), "last_name": excel_text(last),
            "birth_date": as_date(pick(row, headers, "Fecha de nacimiento", "birth_date")) or "",
            "birth_city": excel_text(pick(row, headers, "Ciudad de nacimiento", "Lugar de nacimiento", "birth_city")),
            "birth_country": excel_text(pick(row, headers, "País de nacimiento", "Pais de nacimiento", "birth_country")),
            "nationality": excel_text(pick(row, headers, "Nacionalidad", "nationality")),
            "diocese": excel_text(pick(row, headers, "Diócesis de procedencia", "Diocesis", "diocese")),
            "parish": excel_text(pick(row, headers, "Parroquia de procedencia", "Parroquia", "parish")),
            "address": excel_text(pick(row, headers, "Dirección", "Direccion", "address")),
            "neighborhood": excel_text(pick(row, headers, "Barrio o sector", "Barrio", "Sector", "neighborhood")),
            "city": excel_text(pick(row, headers, "Ciudad de residencia", "Ciudad", "city")),
            "province": excel_text(pick(row, headers, "Provincia o estado", "Provincia", "province")),
            "country": excel_text(pick(row, headers, "País de residencia", "Pais", "country")),
            "personal_phone": excel_text(pick(row, headers, "Teléfono personal", "Celular", "personal_phone"), True),
            "email": excel_text(pick(row, headers, "Correo electrónico", "Correo", "Email", "email")),
            "formation_year": excel_text(pick(row, headers, "Año de formación", "Curso", "formation_year")),
            "room": excel_text(pick(row, headers, "Habitación", "room")),
            "blood_type": excel_text(pick(row, headers, "Tipo de sangre", "Grupo sanguíneo", "blood_type")),
            "allergies": excel_text(pick(row, headers, "Alergias", "allergies")),
            "important_conditions": excel_text(pick(row, headers, "Enfermedades o condiciones médicas", "Condiciones importantes", "important_conditions")),
            "permanent_medicines": excel_text(pick(row, headers, "Medicamentos de uso permanente", "Medicamentos permanentes", "permanent_medicines")),
            "emergency_contact": excel_text(pick(row, headers, "Nombre del contacto de emergencia", "Contacto de emergencia", "emergency_contact")),
            "emergency_relation": excel_text(pick(row, headers, "Parentesco del contacto", "Parentesco", "emergency_relation")),
            "emergency_phone": excel_text(pick(row, headers, "Teléfono del contacto de emergencia", "Teléfono emergencia", "emergency_phone"), True),
            "health_center": excel_text(pick(row, headers, "Centro de salud preferente", "Centro de salud", "health_center")),
            "insurance": excel_text(pick(row, headers, "Seguro médico", "insurance")),
            "dietary_restrictions": excel_text(pick(row, headers, "Restricciones alimentarias", "dietary_restrictions")),
            "notes": excel_text(pick(row, headers, "Observaciones", "Notas", "notes")),
            "consent": excel_text(pick(row, headers, "Autorizo el tratamiento responsable de estos datos", "Consentimiento", "consent")),
        }
        try:
            _, created = db.upsert_seminarian(data)
            result["created" if created else "updated"] += 1
        except ValueError as exc:
            if str(exc).startswith("REVISAR_DUPLICADO:"):
                result["review"] += 1
            else:
                result["errors"] += 1
        except Exception:
            result["errors"] += 1
    return result


def import_inventory(db: Database, filename: str | Path) -> dict[str, int]:
    wb = load_workbook(filename, data_only=True, read_only=True)
    ws = wb["Inventario"] if "Inventario" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = header_map(ws)
    result = {"created": 0, "updated": 0, "recognized": 0, "stock_entries": 0,
              "expired_rejected": 0, "expired_units": 0, "skipped": 0, "errors": 0}
    grouped: dict[str, dict[str, Any]] = {}

    def as_nonnegative_int(value, default=0):
        if value in (None, ""):return default
        number=int(float(str(value).replace(",", ".")))
        if number < 0:raise ValueError("La cantidad no puede ser negativa.")
        return number

    for row in ws.iter_rows(min_row=2, values_only=True):
        name = pick(row, headers, "Producto", "Nombre", "Medicamento")
        if not str(name or "").strip():
            result["skipped"] += 1
            continue
        try:
            result["recognized"]+=1
            stock=as_nonnegative_int(pick(row, headers, "Stock actual", "Cantidad", "Stock", default=0))
            minimum=as_nonnegative_int(pick(row, headers, "Stock mínimo", "Minimo", default=5),5)
            raw_expiry=pick(row,headers,"Fecha de vencimiento","Fecha vencimiento","Fecha_Vencimiento","Vence")
            expiry=as_date(raw_expiry)
            if raw_expiry not in (None,"") and expiry is None:
                raise ValueError("Fecha de vencimiento no reconocida.")
            if expiry and date.fromisoformat(expiry)<date.today():
                result["expired_rejected"]+=1;result["expired_units"]+=stock
                continue
            item_type=normalize(pick(row, headers, "Tipo", default="medicine"))
            if item_type in ("medicamento", "medicina", "medicine"):item_type="medicine"
            elif item_type in ("insumo", "supply"):item_type="supply"
            elif item_type in ("equipo", "equipment"):item_type="equipment"
            else:item_type="medicine"
            category=str(pick(row,headers,"Categoría","Categoria") or "").strip()
            location=str(pick(row,headers,"Ubicación","Ubicacion") or "Por asignar").strip()
            if location.upper()=="QR":location="CU"
            if normalize(location) in {"","por_asignar"}:location=infer_location(category,name)
            data={
                "name":normalize_product_name(name),"type":item_type,
                "generic_name":str(pick(row,headers,"Genérico","Principio activo") or "").strip(),
                "concentration":str(pick(row,headers,"Concentración") or "").strip(),
                "presentation":str(pick(row,headers,"Presentación") or "").strip(),
                "category":category,
                "location":location,
                "unit":str(pick(row,headers,"Unidad",default="unidad") or "unidad").strip(),
                "minimum_stock":minimum,
                "usage":str(pick(row,headers,"Uso","Indicación") or "").strip(),
                "source":str(pick(row,headers,"Fuente","URL") or "").strip(),
                "notes":str(pick(row,headers,"Observaciones","Notas") or "").strip(),
            }
            key=data["name"].casefold();record=grouped.get(key)
            if record is None:
                record={"data":data,"stocks":[]};grouped[key]=record
            else:
                # Mantiene la primera descripción completa y solo rellena campos vacíos.
                for field,value in data.items():
                    if field not in {"name","minimum_stock"} and not record["data"].get(field) and value:record["data"][field]=value
            record["stocks"].append({"quantity":stock,"lot_number":str(pick(row,headers,"Lote") or "").strip(),
                                     "expires_on":expiry})
        except Exception:
            result["errors"] += 1
    batch=db.import_inventory_batch(list(grouped.values()))
    result.update(batch)
    return result


def _style_and_save(wb: Workbook, filename: str | Path):
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123B5D")
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            length = min(42, max(len(str(c.value or "")) for c in column) + 2)
            sheet.column_dimensions[column[0].column_letter].width = length
    wb.save(filename)


def _add_inventory_sheet(wb: Workbook, db: Database, active=False):
    ws = wb.active if active else wb.create_sheet()
    ws.title = "Inventario"
    headers = ["ID","Tipo","Categoría","Producto","Genérico","Concentración","Presentación","Ubicación",
               "Unidad","Stock actual","En cuarentena","Stock mínimo","Próximo vencimiento","Uso","Fuente","Observaciones"]
    ws.append(headers)
    for r in db.inventory():
        ws.append([r["id"],r["type"],r["category"],r["name"],r["generic_name"],r["concentration"],
                   r["presentation"],r["location"],r["unit"],r["stock"],r["quarantined_stock"],r["minimum_stock"],
                   r["earliest_expiry"],r["usage"],r["source"],r["notes"]])


def _add_seminarians_sheet(wb: Workbook, db: Database, active=False):
    sem = wb.active if active else wb.create_sheet()
    sem.title = "Expedientes"
    sem_headers = ["ID","Identificación","Nombres","Apellidos","Fecha de nacimiento","Ciudad de nacimiento","País de nacimiento","Nacionalidad",
                   "Diócesis","Parroquia","Dirección","Barrio o sector","Ciudad de residencia","Provincia o estado","País de residencia","Teléfono personal","Correo electrónico","Año de formación","Habitación",
                   "Tipo de sangre","Alergias","Condiciones importantes","Medicamentos permanentes","Contacto emergencia",
                   "Parentesco","Teléfono emergencia","Centro de salud","Seguro","Restricciones alimentarias","Observaciones"]
    sem.append(sem_headers)
    for r in db.query("SELECT * FROM seminarians WHERE archived=0 ORDER BY last_name,first_name"):
        sem.append([r["id"],r["external_id"],r["first_name"],r["last_name"],r["birth_date"],r["birth_city"],r["birth_country"],r["nationality"],
                    r["diocese"],r["parish"],r["address"],r["neighborhood"],r["city"],r["province"],r["country"],r["personal_phone"],r["email"],r["formation_year"],r["room"],
                    r["blood_type"],r["allergies"],r["important_conditions"],r["permanent_medicines"],r["emergency_contact"],
                    r["emergency_relation"],r["emergency_phone"],r["health_center"],r["insurance"],r["dietary_restrictions"],r["notes"]])


def _add_movements_sheet(wb: Workbook, db: Database, active=False):
    mov = wb.active if active else wb.create_sheet()
    mov.title = "Movimientos"
    mov.append(["ID","Fecha","Tipo","Producto","Seminarista","Cantidad","Dosis","Motivo","Estado","Motivo anulación"])
    for r in db.query("""SELECT m.id,m.occurred_at,m.type,i.name product,
        CASE WHEN s.id IS NOT NULL THEN s.first_name||' '||s.last_name WHEN m.type IN ('delivery','administered','loan','return') THEN 'Expediente eliminado' ELSE '' END seminarian,m.quantity,m.dose,m.reason,
        CASE WHEN m.voided_at IS NULL THEN 'Vigente' ELSE 'Anulado' END state,m.void_reason
        FROM movements m JOIN items i ON i.id=m.item_id LEFT JOIN seminarians s ON s.id=m.seminarian_id
        ORDER BY m.occurred_at DESC"""):
        mov.append(list(r))


def _add_medical_controls_sheet(wb: Workbook, db: Database, active=False, seminarian_id: int | None = None):
    ws=wb.active if active else wb.create_sheet();ws.title="Controles de salud"
    ws.append(["ID","Seminarista","Fecha","Peso kg","Estatura cm","IMC","Sistólica","Diastólica","Pulso",
               "Glucosa mg/dL","Condición glucosa","SpO2 %","Temperatura °C","Observaciones","Recomendación o derivación","Registrado por","Estado","Motivo anulación"])
    query="""SELECT c.id,s.first_name||' '||s.last_name seminarian,c.measured_at,c.weight_kg,c.height_cm,c.bmi,
        c.systolic,c.diastolic,c.pulse,c.glucose_mg_dl,c.glucose_context,c.oxygen_saturation,c.temperature_c,
        c.notes,c.referral,c.recorded_by,CASE WHEN c.voided_at IS NULL THEN 'Vigente' ELSE 'Anulado' END state,c.void_reason
        FROM medical_controls c JOIN seminarians s ON s.id=c.seminarian_id"""
    params=()
    if seminarian_id is not None:query+=" WHERE c.seminarian_id=? AND c.voided_at IS NULL";params=(seminarian_id,)
    else:query+=" WHERE c.voided_at IS NULL"
    query+=" ORDER BY c.measured_at DESC,c.id DESC"
    for r in db.query(query,params):
        ws.append(list(r))


def export_inventory(db: Database, filename: str | Path):
    wb=Workbook(); _add_inventory_sheet(wb,db,True); _style_and_save(wb,filename)


def export_seminarians(db: Database, filename: str | Path):
    wb=Workbook();_add_seminarians_sheet(wb,db,True);_add_medical_controls_sheet(wb,db);_style_and_save(wb,filename)


def export_medical_controls(db: Database, filename: str | Path, seminarian_id: int | None = None):
    wb=Workbook();_add_medical_controls_sheet(wb,db,True,seminarian_id);_style_and_save(wb,filename)


def export_movements(db: Database, filename: str | Path):
    wb=Workbook(); _add_movements_sheet(wb,db,True); _style_and_save(wb,filename)


def export_all(db: Database, filename: str | Path):
    wb = Workbook()
    _add_inventory_sheet(wb,db,True);_add_seminarians_sheet(wb,db);_add_medical_controls_sheet(wb,db);_add_movements_sheet(wb,db)
    _style_and_save(wb,filename)
